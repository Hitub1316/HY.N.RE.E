"""
parsers.py — Multi-format novel text parser (.txt, .pdf, .docx).

Supports:
  - parse_txt: Plain text with UTF-8 / Latin-1 decoding.
  - parse_pdf: PDF text extraction via pypdf with header/footer stripping.
  - parse_docx: Word document extraction via python-docx.
  - parse_file: Unified dispatcher by filename extension.
"""

import io
import re
from pathlib import Path
from typing import Union


def parse_txt(file_input: Union[bytes, str, Path]) -> str:
    """Extract string content from .txt input (bytes or path)."""
    if isinstance(file_input, (str, Path)):
        path = Path(file_input)
        raw = path.read_bytes()
    else:
        raw = file_input

    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        return raw.decode("latin-1")


def _strip_pdf_headers_footers(page_texts: list[str]) -> str:
    """
    Remove running headers, page numbers, and repeating margin lines across PDF pages.
    """
    if not page_texts:
        return ""

    # Tokenize each page into lines
    lines_per_page = [[line.strip() for line in pt.splitlines() if line.strip()] for pt in page_texts]
    
    # Track line occurrences at top 2 and bottom 2 positions of pages
    margin_line_counts = {}
    total_pages = len(page_texts)

    for page_lines in lines_per_page:
        if not page_lines:
            continue
        margins = set()
        # Top 2 lines
        for line in page_lines[:2]:
            margins.add(line)
        # Bottom 2 lines
        for line in page_lines[-2:]:
            margins.add(line)
        
        for line in margins:
            margin_line_counts[line] = margin_line_counts.get(line, 0) + 1

    # Lines appearing in top/bottom margins of >= 30% of pages are candidate headers/footers
    repeating_margins = {line for line, count in margin_line_counts.items() if count >= max(2, total_pages * 0.3)}

    cleaned_pages = []
    for page_lines in lines_per_page:
        filtered_lines = []
        for line in page_lines:
            # Check for standalone page digits ("12", "Page 12", "Page 12 of 300")
            if re.match(r"^(page\s+)?\d+(\s+of\s+\d+)?$", line, re.IGNORECASE):
                continue
            # Check if it matches a repeated margin header/footer
            if line in repeating_margins:
                continue
            filtered_lines.append(line)
        cleaned_pages.append("\n".join(filtered_lines))

    return "\n\n".join(cleaned_pages)


def parse_pdf(file_input: Union[bytes, str, Path]) -> str:
    """Extract clean text from .pdf input via pypdf with header/footer removal."""
    import pypdf

    if isinstance(file_input, (str, Path)):
        reader = pypdf.PdfReader(str(file_input))
    else:
        reader = pypdf.PdfReader(io.BytesIO(file_input))

    page_texts = []
    for page in reader.pages:
        txt = page.extract_text() or ""
        if txt.strip():
            page_texts.append(txt)

    return _strip_pdf_headers_footers(page_texts)


def parse_docx(file_input: Union[bytes, str, Path]) -> str:
    """Extract clean text from .docx input via python-docx."""
    import docx

    if isinstance(file_input, (str, Path)):
        doc = docx.Document(str(file_input))
    else:
        doc = docx.Document(io.BytesIO(file_input))

    paragraphs = []
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            paragraphs.append(txt)

    return "\n\n".join(paragraphs)


def parse_file(file_name: str, file_input: Union[bytes, str, Path]) -> str:
    """
    Dispatch parsing by file extension (.txt, .pdf, .docx).
    Raises ValueError for unsupported formats.
    """
    ext = Path(file_name).suffix.lower()
    if ext in (".txt", ".text"):
        return parse_txt(file_input)
    elif ext == ".pdf":
        return parse_pdf(file_input)
    elif ext == ".docx":
        return parse_docx(file_input)
    else:
        raise ValueError(f"Unsupported file extension '{ext}'. Supported formats: .txt, .pdf, .docx")
